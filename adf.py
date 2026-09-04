"""
Asylum Document Filler

Tornado web app for filling DOCX templates and optionally converting them to PDF
with LibreOffice.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import tornado.ioloop
import tornado.web
from docx import Document


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = Path(r"\\IT-PC2\Scan_Clients\ASYLUM Letters")
TEMPLATE_DIR = BASE_DIR
WEB_TEMPLATE_DIR = BASE_DIR / "templates"
OUTPUT_DIR.mkdir(exist_ok=True)


@dataclass(frozen=True)
class DocumentKind:
    key: str
    label: str
    template_name: str

    @property
    def template_path(self) -> Path:
        return TEMPLATE_DIR / self.template_name

    @property
    def output_stem(self) -> str:
        return self.template_path.stem.replace("_Template", "")


DOCUMENT_KINDS = {
    "newcomer": DocumentKind("newcomer", "Asylum Newcomer", "ASYLUM_NEW_COMER_Template - adjusted for Concourt judgement 03 september.docx"),
    "renewal": DocumentKind("renewal", "Asylum Renewal", "ASYLUM_RENEWAL_Template.docx"),
    "confirmation": DocumentKind("confirmation", "Confirmation Letter", "CONFIRMATION_LETTER_Template - Copy.docx"),
    "explainer": DocumentKind("explainer", "Explainer", "Explainer_CONFIRMATION_LETTER_Template - New.docx")
}


TITLE_WORDS = {
    "Mr": {"pronoun": "his", "gender": "He", "object": "him"},
    "Ms": {"pronoun": "her", "gender": "She", "object": "her"},
    "Mrs": {"pronoun": "her", "gender": "She", "object": "her"},
}

DEFAULT_WORDS = {"pronoun": "their", "gender": "They", "object": "them"}


def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "", value).strip()
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned or "Applicant"


def _clear_paragraph_runs(paragraph) -> None:
    """Remove all runs from a paragraph without relying on Paragraph.clear(),
    which is not available on every python-docx version."""
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    for placeholder, value in replacements.items():

        if placeholder not in paragraph.text:
            continue

        # ----- Special formatting for identity placeholders -----

        if placeholder in ("[Identity Block]", "[Identity Sentence]"):

            # Preserve text before/after the placeholder
            before, after = paragraph.text.split(placeholder, 1)

            _clear_paragraph_runs(paragraph)

            if before:
                paragraph.add_run(before)

            run = paragraph.add_run(value)
            run.bold = True
            run.underline = True

            if after:
                paragraph.add_run(after)

            # NOTE: continue (not return) - a paragraph can contain more than
            # one placeholder and earlier code was bailing out of the whole
            # function here, silently skipping every other placeholder left
            # in this paragraph.
            continue

        # ----- Normal replacements -----

        full = "".join(run.text for run in paragraph.runs)
        full = full.replace(placeholder, value)

        if paragraph.runs:
            paragraph.runs[0].text = full
            for run in paragraph.runs[1:]:
                run.text = ""


def replace_in_table(table, replacements: dict[str, str]) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(paragraph, replacements)
            for nested_table in cell.tables:
                replace_in_table(nested_table, replacements)


def replace_everywhere(document, replacements: dict[str, str]) -> None:
    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph, replacements)

    for table in document.tables:
        replace_in_table(table, replacements)

    for section in document.sections:
        header_footers = (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        )
        for part in header_footers:
            for paragraph in part.paragraphs:
                replace_in_paragraph(paragraph, replacements)
            for table in part.tables:
                replace_in_table(table, replacements)


def libreoffice_path() -> str | None:
    configured = os.environ.get("LIBREOFFICE_PATH")
    candidates = [
        configured,
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    return None


def convert_to_pdf(docx_path: Path) -> Path | None:
    """Best-effort DOCX -> PDF conversion. Returns None (instead of raising)
    when LibreOffice isn't installed/found, matching the "PDFs appear too
    when LibreOffice is available" behaviour described in the UI."""
    libreoffice = libreoffice_path()
    if not libreoffice:
        return None

    try:
        subprocess.run(
            [
                libreoffice,
                "--headless",
                "--convert-to",
                "pdf",
                str(docx_path),
                "--outdir",
                str(docx_path.parent),
            ],
            check=True,
            timeout=60,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
        return None

    pdf_path = docx_path.with_suffix(".pdf")
    return pdf_path if pdf_path.exists() else None


def generate_document(
    kind: DocumentKind,
    title: str,
    full_names: str,
    dob: str,
    tracking_number: str = "",
    nationality: str = "",
    application_type: str = "",
    passport_number: str = "",
    selected_gender: str = "",
    last_name: str = "",
    identinty_block: str = "",
    identinty_sentence: str = "",
) -> tuple[Path, Path | None]:
    words = TITLE_WORDS.get(title, DEFAULT_WORDS)
    if selected_gender == "male":
        words = {"pronoun": "his", "gender": "He", "object": "him"}
    elif selected_gender == "female":
        words = {"pronoun": "her", "gender": "She", "object": "her"}

    # ---------------- Confirmation identity ----------------

    if passport_number:
        identity_block = f"Date of birth: {dob}\n\nPassport No. {passport_number}"
        identity_sentence = f"{full_names}, Passport No. {passport_number}"
    else:
        identity_block = f"Date of birth: {dob}"
        identity_sentence = f"{full_names}, Date of birth: {dob}"
# date shoulde be fixed now 11 aug
    today = datetime.today().strftime("%d %B %Y")
    replacements = {
        "[Todays Date]": today,
        "[Today's Date]": today,
        "[Date]": today,
        "[title]": title,
        "[Title]": title,
        "[full names]": full_names,
        "[Full Names]": full_names,
        "[lastName]": last_name,
        "[Applicant Name]": full_names,
        "[DOB]": dob,
        "[dob]": dob,
        "[Date of Birth]": dob,
        "[pronoun]": words["pronoun"],
        "[gender]": words["gender"],
        "[tracking number]": tracking_number,
        "[Tracking Number]": tracking_number,
        "[Nationality]": nationality,
        "[nationality]": nationality,
        "[Application type]": application_type,
        "[Application Type]": application_type,
        "[Passport Number]": passport_number,
        "[Passport No]": passport_number,
        "[passport number]": passport_number,
        "[Identity Block]": identity_block,
        "[Identity Sentence]": identity_sentence,
    }

    if kind.key == "renewal":
        replacements.update({
            "[Todays date]": today,
            "[Applicant Name]": full_names,
            "[tracking number]": tracking_number,
            "[Nationality]": nationality,
        })

    if kind.key == "confirmation":
        replacements.update({
            "[Applicant Name]": full_names,
            "[Applicant date of birth]": dob,
            "[Passport Number]": passport_number,
            "[Application type]": application_type or "Asylum Seekers",
            "[him/her]": words["gender"],
            "[Identity Block]": identity_block,
            "[Identity Sentence]": identity_sentence,
        })

    document = Document(kind.template_path)
    replace_everywhere(document, replacements)

    output_path = OUTPUT_DIR / f"{kind.output_stem}_FOR_{safe_filename(full_names)}.docx"
    document.save(output_path)
    return output_path, convert_to_pdf(output_path)


class BaseHandler(tornado.web.RequestHandler):
    @property
    def kinds(self) -> dict[str, DocumentKind]:
        return self.application.settings["document_kinds"]


class DashboardHandler(BaseHandler):
    def get(self, kind_key: str = "newcomer") -> None:
        kind = self.kinds.get(kind_key)
        if not kind:
            raise tornado.web.HTTPError(404)

        self.render(
            "index.html",
            kinds=self.kinds,
            current_kind=kind,
            template_ok=kind.template_path.exists(),
        )


class GenerateHandler(BaseHandler):
    def post(self, kind_key: str) -> None:
        kind = self.kinds.get(kind_key)
        if not kind:
            self.set_status(404)
            self.write({"ok": False, "error": "Unknown document type."})
            return

        if not kind.template_path.exists():
            self.set_status(400)
            self.write({"ok": False, "error": f"Template not found: {kind.template_name}"})
            return

        title = self.get_body_argument("title", "").strip()
        full_names = self.get_body_argument("full_names", "").strip()
        dob = self.get_body_argument("dob", "").strip()
        nationality = self.get_body_argument("nationality", "").strip()
        tracking_number = self.get_body_argument("tracking_number", "").strip()
        application_type = self.get_body_argument("application_type", "").strip()
        passport_number = self.get_body_argument("passport_number", "").strip()
        selected_gender = self.get_body_argument("gender", "").strip().lower()
        last_name = self.get_body_argument("last_name", "").strip()

        if not application_type:
            application_type = "Asylum Seekers"
        
        if not full_names:
            self.set_status(400)
            self.write({"ok": False, "error": "Full names are required."})
            return

        try:
            docx_path, pdf_path = generate_document(
                kind,
                title,
                full_names,
                dob,
                tracking_number,
                nationality,
                application_type,
                passport_number,
                selected_gender,
                last_name,
            )
        except Exception as exc:
            self.set_status(500)
            self.write({"ok": False, "error": str(exc)})
            return

        self.write(
            {
                "ok": True,
                "filename": docx_path.name,
                "pdf": pdf_path.name if pdf_path else None,
                "pdf_available": pdf_path is not None,
            }
        )


class FilesHandler(BaseHandler):
    def get(self, kind_key: str) -> None:
        kind = self.kinds.get(kind_key)
        if not kind:
            self.set_status(404)
            self.set_header("Content-Type", "application/json")
            self.write(json.dumps([]))
            return

        files = sorted(
            OUTPUT_DIR.glob(f"{kind.output_stem}_FOR_*.docx"),
            key=os.path.getmtime,
            reverse=True,
        )
        result = []
        for file_path in files:
            pdf_path = file_path.with_suffix(".pdf")
            stat = file_path.stat()
            result.append(
                {
                    "name": file_path.name,
                    "pdf": pdf_path.name if pdf_path.exists() else None,
                    "size_kb": round(stat.st_size / 1024, 1),
                    "created": datetime.fromtimestamp(stat.st_mtime).strftime("%d %b %Y %H:%M"),
                }
            )

        self.set_header("Content-Type", "application/json")
        self.write(json.dumps(result))


class PdfHandler(tornado.web.RequestHandler):
    def get(self, filename: str) -> None:
        filepath = (OUTPUT_DIR / filename).resolve()
        if OUTPUT_DIR.resolve() not in filepath.parents or not filepath.exists():
            self.set_status(404)
            return

        self.set_header("Content-Type", "application/pdf")
        with open(filepath, "rb") as f:
            self.write(f.read())


class DownloadHandler(tornado.web.RequestHandler):
    def get(self, filename: str) -> None:
        path = (OUTPUT_DIR / filename).resolve()
        if OUTPUT_DIR.resolve() not in path.parents or path.suffix.lower() not in {".docx", ".pdf"} or not path.exists():
            raise tornado.web.HTTPError(404)

        if path.suffix.lower() == ".pdf":
            content_type = "application/pdf"
        else:
            content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        self.set_header("Content-Type", content_type)
        self.set_header("Content-Disposition", f'attachment; filename="{path.name}"')
        with path.open("rb") as file_handle:
            self.write(file_handle.read())


class DeleteHandler(tornado.web.RequestHandler):
    def post(self, filename: str) -> None:
        filepath = (OUTPUT_DIR / filename).resolve()
        if OUTPUT_DIR.resolve() not in filepath.parents or filepath.suffix != ".docx" or not filepath.exists():
            self.set_status(404)
            self.write({"ok": False, "error": "File not found."})
            return

        filepath.unlink()
        pdf_path = filepath.with_suffix(".pdf")
        if pdf_path.exists():
            pdf_path.unlink()

        self.write({"ok": True})


def make_app() -> tornado.web.Application:
    return tornado.web.Application(
        [
            (r"/", DashboardHandler),
            (r"/document/([a-z-]+)", DashboardHandler),
            (r"/generate/([a-z-]+)", GenerateHandler),
            (r"/files/([a-z-]+)", FilesHandler),
            (r"/download/([^/]+\.(?:docx|pdf))", DownloadHandler),
            (r"/delete/([^/]+\.docx)", DeleteHandler),
            (r"/pdf/([^/]+\.pdf)", PdfHandler),
        ],
        document_kinds=DOCUMENT_KINDS,
        template_path=str(WEB_TEMPLATE_DIR),
        debug=True,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Asylum Document Filler")
    parser.add_argument("--port", type=int, default=2026)
    args = parser.parse_args()

    app = make_app()
    app.listen(args.port)
    # prints
    print("\nAsylum Document Filler")
    print(f"URL       -> http://localhost:{args.port}")
    print(f"Templates -> {TEMPLATE_DIR}")
    print(f"Output    -> {OUTPUT_DIR}\n")

    tornado.ioloop.IOLoop.current().start()


if __name__ == "__main__":
    main()
